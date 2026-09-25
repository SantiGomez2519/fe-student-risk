"""
Carga de datos desde el Excel del formulario (Form Responses 1) y
cálculo del resultado de riesgo por estudiante.

Reutilizable desde:
  - management command `load_data`
  - endpoint de subida de Excel
"""

from __future__ import annotations

import openpyxl
from django.utils import timezone

from .models import RiskResult, Student
from .scoring import score_response

# Mapeo columna del formulario -> campo del modelo Student (caracterización)
CHARACTERIZATION = {
    "Nombre Completo": "full_name",
    "Edad": "age",
    "Sexo": "sex",
    "Programa académico": "program",
    "Semestre actual": "semester",
    "Ciudad de procedencia": "city_origin",
    "Ciudad donde vive actualmente": "city_current",
    "Actualmente vive con…": "lives_with",
    "¿Es la primera persona de su familia en ingresar a la universidad?": "first_generation",
}

RESPONSES_SHEET = "Form Responses 1"


def _clean(value):
    if value is None:
        return None
    if isinstance(value, str):
        return value.strip()
    return value


def read_rows(path_or_file, sheet_name: str | None = None):
    """Lee el Excel y devuelve (headers, list[dict fila]) de la hoja de respuestas."""
    wb = openpyxl.load_workbook(path_or_file, data_only=True)
    name = sheet_name or (RESPONSES_SHEET if RESPONSES_SHEET in wb.sheetnames else wb.sheetnames[0])
    ws = wb[name]
    headers = [_clean(ws.cell(row=1, column=c).value) for c in range(1, ws.max_column + 1)]
    rows = []
    for r in range(2, ws.max_row + 1):
        row = {}
        empty = True
        for c, header in enumerate(headers, start=1):
            if not header:
                continue
            v = _clean(ws.cell(row=r, column=c).value)
            row[header] = v
            if v not in (None, ""):
                empty = False
        if not empty:
            rows.append(row)
    wb.close()
    return headers, rows


def _to_int(v):
    try:
        return int(float(v))
    except (TypeError, ValueError):
        return None


def upsert_student(row: dict) -> tuple[Student, bool]:
    """Crea o actualiza un estudiante + su resultado a partir de una fila."""
    name = row.get("Nombre Completo") or row.get("Nombre completo")
    if not name:
        # Sin nombre no podemos identificar de forma estable; usamos timestamp.
        name = f"(sin nombre) {row.get('Timestamp')}"

    defaults = {"raw": {k: (v.isoformat() if hasattr(v, 'isoformat') else v)
                         for k, v in row.items()}}
    for header, field in CHARACTERIZATION.items():
        if header == "Nombre Completo":
            continue
        val = row.get(header)
        if field == "age":
            val = _to_int(val)
        elif field == "semester":
            iv = _to_int(val)
            val = str(iv) if iv is not None else (str(val) if val else "")
        else:
            val = val or ""
        defaults[field] = val
    ts = row.get("Timestamp")
    if hasattr(ts, "isoformat"):
        if timezone.is_naive(ts):
            ts = timezone.make_aware(ts, timezone.get_current_timezone())
        defaults["submitted_at"] = ts
    else:
        defaults["submitted_at"] = None

    student, created = Student.objects.update_or_create(
        full_name=str(name).strip(), defaults=defaults
    )

    analysis = score_response(row)
    RiskResult.objects.update_or_create(
        student=student,
        defaults={
            "total": analysis["total"],
            "total_max": analysis["total_max"],
            "total_pct": analysis["total_pct"],
            "level": analysis["level"],
            "dimensions": analysis["dimensions"],
            "items": analysis["items"],
            "alerts": analysis["alerts"],
        },
    )
    return student, created


def load_from_path(path, sheet_name: str | None = None) -> dict:
    _, rows = read_rows(path, sheet_name)
    created = updated = 0
    for row in rows:
        _, was_created = upsert_student(row)
        if was_created:
            created += 1
        else:
            updated += 1
    return {"rows": len(rows), "created": created, "updated": updated}
